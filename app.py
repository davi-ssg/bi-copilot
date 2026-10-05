import streamlit as st
import pandas as pd
import plotly.express as px
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from database import (
    init_db, 
    get_database_schema, 
    load_df_to_sqlite, 
    reset_database, 
    run_query
)
from agent import run_bi_agent

# 1. Configuração da página e inicialização do banco
st.set_page_config(page_title="BI Copilot", page_icon="📊", layout="wide")
init_db()

# 2. Gestão do Session State
if "messages" not in st.session_state:
    st.session_state.messages = []

if "uploader_key" not in st.session_state:
    st.session_state.uploader_key = 0

schema_atual = get_database_schema()
has_data = "-- Tabela:" in schema_atual

# ==========================================
# FUNÇÕES AUXILIARES (GRÁFICOS E EXPORTAÇÃO)
# ==========================================
def convert_to_docx(response_text: str, df: pd.DataFrame = None) -> bytes:
    """Gera um relatório profissional em formato Word (.docx)."""
    doc = Document()

    # Estilo Global / Título
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = title_p.add_run("📊 Relatório Executivo de BI")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(31, 78, 121) # Azul corporativo

    sub_p = doc.add_paragraph()
    run_sub = sub_p.add_run("Gerado automaticamente pelo BI Copilot")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph() # Espaçador

    # Seção 1: Análise Sintetizada
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Análise e Insights")
    run_h1.font.name = 'Arial'
    run_h1.font.color.rgb = RGBColor(31, 78, 121)

    # Insere o texto da resposta linha a linha
    for line in response_text.split("\n"):
        line_clean = line.strip()
        if line_clean:
            p = doc.add_paragraph()
            run = p.add_run(line_clean)
            run.font.name = 'Arial'
            run.font.size = Pt(11)

    # Seção 2: Tabela de Dados Formatada (se existir)
    if df is not None and not df.empty:
        doc.add_paragraph()
        h2 = doc.add_heading(level=1)
        run_h2 = h2.add_run("2. Detalhamento dos Dados")
        run_h2.font.name = 'Arial'
        run_h2.font.color.rgb = RGBColor(31, 78, 121)

        # Criar tabela estilizada
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = 'Table Grid'

        # Cabeçalho da tabela
        hdr_cells = table.rows[0].cells
        for idx, col_name in enumerate(df.columns):
            hdr_cells[idx].text = str(col_name).replace('_', ' ').title()
            # Deixar texto do cabeçalho em negrito
            for p in hdr_cells[idx].paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.name = 'Arial'

        # Linhas de dados (limite de 100 linhas no Word para performance)
        for _, row in df.head(100).iterrows():
            row_cells = table.add_row().cells
            for idx, val in enumerate(row):
                text_val = str(val) if pd.notna(val) else ""
                row_cells[idx].text = text_val
                for p in row_cells[idx].paragraphs:
                    for r in p.runs:
                        r.font.name = 'Arial'
                        r.font.size = Pt(10)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_export_options(response_text: str, df: pd.DataFrame, key_suffix: str):
    """Exibe os botões de download de dados e relatórios."""
    st.markdown("##### 📥 Exportar Resultados")
    col_doc, col_csv, col_excel = st.columns(3)
    
    with col_doc:
        try:
            docx_bytes = convert_to_docx(response_text, df)
            st.download_button(
                label="📝 Baixar Relatório (Word)",
                data=docx_bytes,
                file_name="relatorio_executivo.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"dl_docx_{key_suffix}",
                use_container_width=True
            )
        except Exception as e:
            st.error(f"Erro ao gerar Word: {e}")

    if df is not None and not df.empty:
        with col_csv:
            csv_bytes = df.to_csv(index=False).encode('utf-8')
            st.download_button(
                label="📊 Baixar Tabela (CSV)",
                data=csv_bytes,
                file_name="dados_consulta.csv",
                mime="text/csv",
                key=f"dl_csv_{key_suffix}",
                use_container_width=True
            )
            
        with col_excel:
            try:
                excel_bytes = convert_df_to_excel(df)
                st.download_button(
                    label="📗 Baixar Tabela (Excel)",
                    data=excel_bytes,
                    file_name="dados_consulta.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"dl_xlsx_{key_suffix}",
                    use_container_width=True
                )
            except Exception:
                pass

def render_chart(df: pd.DataFrame):
    """Gera gráficos interativos usando Plotly Express."""
    if df is None or df.empty or len(df.columns) < 2:
        return

    try:
        temp_df = df.copy()
        cols = list(temp_df.columns)
        numeric_cols = temp_df.select_dtypes(include=['number']).columns.tolist()
        
        if not numeric_cols:
            return

        x_col = cols[0]
        y_col = numeric_cols[0]
        title_text = f"<b>{y_col.replace('_', ' ').title()}</b> por <b>{x_col.replace('_', ' ').title()}</b>"

        is_time = any(kw in str(x_col).lower() for kw in ['data', 'date', 'mes', 'ano', 'dia', 'tempo', 'periodo'])

        st.markdown("---")

        if is_time:
            fig = px.line(
                temp_df, x=x_col, y=y_col, markers=True,
                title=title_text, template="plotly_dark"
            )
            fig.update_traces(line_color="#6366F1", line_width=3, marker_size=8)
        elif len(temp_df) <= 5 and len(cols) == 2:
            fig = px.pie(
                temp_df, names=x_col, values=y_col, hole=0.45,
                title=title_text, template="plotly_dark",
                color_discrete_sequence=px.colors.qualitative.Pastel
            )
            fig.update_traces(textposition='inside', textinfo='percent+label')
        else:
            fig = px.bar(
                temp_df, x=x_col, y=y_col, text_auto=True,
                title=title_text, template="plotly_dark",
                color=y_col, color_continuous_scale="Viridis"
            )
            fig.update_layout(coloraxis_showscale=False)

        fig.update_layout(
            margin=dict(l=20, r=20, t=50, b=20),
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            font=dict(family="Arial, sans-serif", size=13),
            xaxis_title=x_col.replace('_', ' ').title(),
            yaxis_title=y_col.replace('_', ' ').title()
        )

        st.plotly_chart(fig, use_container_width=True)

    except Exception:
        pass

# ==========================================
# 3. BARRA LATERAL (SIDEBAR)
# ==========================================
with st.sidebar:
    st.title("🤖 BI Copilot")
    st.caption("Powered by Gemini API & SQLite")
    st.divider()
    
    st.subheader("📁 Carregar Base de Dados")
    with st.popover("ℹ️ Como formatar a planilha?"):
        st.markdown("""
        - **Cabeçalho na 1ª linha:** Nomes claros de colunas.
        - **Sem células mescladas.**
        - **Formatos:** Números limpos e datas no formato YYYY-MM-DD.
        """)

    uploaded_file = st.file_uploader(
        "Envie um arquivo CSV ou Excel", 
        type=["csv", "xlsx"],
        key=f"uploader_{st.session_state.uploader_key}"
    )
    
    if uploaded_file is not None:
        try:
            if uploaded_file.name.endswith(".csv"):
                user_df = pd.read_csv(uploaded_file)
            else:
                user_df = pd.read_excel(uploaded_file)
                
            load_df_to_sqlite(user_df, table_name="dados_usuario")
            st.success(f" Base carregada com sucesso! ({len(user_df)} linhas)")
            
            with st.expander("👁️ Pré-visualizar Colunas", expanded=True):
                st.dataframe(user_df.head(5), use_container_width=True)
                
            schema_atual = get_database_schema()
            has_data = "-- Tabela:" in schema_atual
        except Exception as e:
            st.error(f"Erro ao processar arquivo: {e}")

    if st.button("🗑️ Zerar / Apagar Banco de Dados", use_container_width=True):
        reset_database()
        st.session_state.messages = []
        st.session_state.uploader_key += 1
        st.success("Banco de dados e histórico limpos!")
        st.rerun()

    st.divider()
    if st.button("💬 Limpar Conversa", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

    st.divider()
    with st.expander("🗄️ Esquema Atual do Banco"):
        st.code(schema_atual, language="sql")

# ==========================================
# 4. TELA INICIAL CENTRALIZADA
# ==========================================
prompt_sugerido = None

if not st.session_state.messages:
    _, center_col, _ = st.columns([1, 2.5, 1])
    with center_col:
        st.markdown("""
            <div style="text-align: center; padding: 2rem 0;">
                <h1 style="font-size: 2.5rem;">📊 BI Copilot</h1>
                <p style="color: #6c757d; font-size: 1.1rem;">Carregue a sua planilha CSV ou Excel e faça análises em linguagem natural.</p>
            </div>
        """, unsafe_allow_html=True)

        if not has_data:
            st.info("👈 **Para começar:** Envie um arquivo CSV ou Excel no menu lateral.")
        else:
            st.write("##### 💡 O que deseja analisar primeiro?")
            c1, c2 = st.columns(2)
            with c1:
                if st.button("📋 Resumo da Base", use_container_width=True):
                    prompt_sugerido = "Faça um resumo geral dos dados, explicando o que cada coluna representa e o total de registros."
                if st.button("📈 Principais Métricas", use_container_width=True):
                    prompt_sugerido = "Quais são as principais métricas numéricas e somatórios presentes nesta base de dados?"
            with c2:
                if st.button("🏆 Principais Destaques", use_container_width=True):
                    prompt_sugerido = "Quais são os itens mais frequentes ou com maiores valores na tabela?"
                if st.button("💡 Sugerir Perguntas", use_container_width=True):
                    prompt_sugerido = "Com base nas colunas desta tabela, sugira 3 perguntas analíticas interessantes que eu posso fazer."

# ==========================================
# 5. HISTÓRICO DE MENSAGENS
# ==========================================
for idx, msg in enumerate(st.session_state.messages):
    with st.chat_message(msg["role"]):
        st.write(msg["content"])
        
        if msg.get("sql_used"):
            with st.expander("🔍 Ver Consulta SQL Executada"):
                st.code(msg["sql_used"], language="sql")
                
        if msg.get("df") is not None and not msg["df"].empty:
            st.dataframe(msg["df"], use_container_width=True)
            render_chart(msg["df"])
            
        if msg["role"] == "assistant":
            render_export_options(msg["content"], msg.get("df"), key_suffix=str(idx))

# ==========================================
# 6. ENTRADA DO USUÁRIO & PROCESSAMENTO
# ==========================================
user_input = prompt_sugerido or st.chat_input("Digite a sua pergunta sobre a base de dados...")

if user_input:
    if not has_data:
        st.warning("⚠️ **Nenhuma base de dados carregada!** Envie um arquivo CSV ou Excel na barra lateral antes de fazer perguntas.")
    else:
        st.session_state.messages.append({"role": "user", "content": user_input})
        
        with st.chat_message("user"):
            st.write(user_input)

        with st.chat_message("assistant"):
            with st.spinner("🤖 O Gemini está a consultar o banco de dados..."):
                result = run_bi_agent(user_input)

            if result["status"] == "success":
                response_text = result["response"]
                sql_used = result.get("sql_used")
                raw_data = result.get("raw_data")
                df = pd.DataFrame(raw_data) if raw_data else None

                st.write(response_text)
                if sql_used:
                    with st.expander("🔍 Ver Consulta SQL Executada"):
                        st.code(sql_used, language="sql")
                if df is not None and not df.empty:
                    st.dataframe(df, use_container_width=True)
                    render_chart(df)

                render_export_options(response_text, df, key_suffix=str(len(st.session_state.messages)))

                st.session_state.messages.append({
                    "role": "assistant",
                    "content": response_text,
                    "sql_used": sql_used,
                    "df": df
                })
            else:
                erro_msg = f"Erro no processamento: {result['response']}"
                st.error(erro_msg)
                st.session_state.messages.append({
                    "role": "assistant",
                    "content": erro_msg,
                    "sql_used": None,
                    "df": None
                })
        
        st.rerun()